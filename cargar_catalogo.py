"""
Carga el catálogo completo desde 'catalogo_completo_real.xlsx'.
Hace upsert: actualiza los productos existentes por referencia e inserta los nuevos.
"""
import os
from pathlib import Path
import openpyxl
from decimal import Decimal
from sqlalchemy.orm import Session
from app.core.models import SessionLocal, Producto, EscalaDescuento, crear_tablas

EXCEL_PATH = Path("data/imports/catalogo_completo_real.xlsx")


def cargar():
    crear_tablas()
    db: Session = SessionLocal()
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        hoja = wb["Productos"]

        # Leer encabezados
        encabezados = [str(c.value).strip().lower() if c.value else "" for c in hoja[1]]
        idx = {nombre: i for i, nombre in enumerate(encabezados)}
        print(f"Columnas detectadas: {encabezados}")

        creados = 0
        actualizados = 0
        errores = []

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if not fila or all(v is None for v in fila):
                continue
            try:
                ref = str(fila[idx["referencia"]]).strip()
                cat = str(fila[idx["categoria"]]).strip()
                desc = str(fila[idx["descripcion"]]).strip()
                precio = float(fila[idx["precio_unitario"]])

                dimension = None
                if "dimension" in idx and fila[idx["dimension"]] is not None:
                    dimension = str(fila[idx["dimension"]]).strip()

                especificaciones = None
                if "especificaciones" in idx and fila[idx["especificaciones"]] is not None:
                    especificaciones = str(fila[idx["especificaciones"]]).strip()

                iva_pct = 19.0
                if "iva_pct" in idx and fila[idx["iva_pct"]] is not None:
                    iva_pct = float(fila[idx["iva_pct"]])

                aplica_escala = False
                if "aplica_descuento_escala" in idx and fila[idx["aplica_descuento_escala"]] is not None:
                    valor = str(fila[idx["aplica_descuento_escala"]]).strip().lower()
                    aplica_escala = valor in ("true", "si", "sí", "1", "yes")

                # Upsert
                existente = db.query(Producto).filter(Producto.referencia == ref).first()
                if existente:
                    existente.categoria = cat
                    existente.descripcion = desc
                    existente.dimension = dimension
                    existente.especificaciones = especificaciones
                    existente.precio_unitario = precio
                    existente.iva_pct = iva_pct
                    existente.aplica_descuento_escala = aplica_escala
                    existente.activo = True
                    actualizados += 1
                else:
                    db.add(Producto(
                        referencia=ref,
                        categoria=cat,
                        descripcion=desc,
                        dimension=dimension,
                        especificaciones=especificaciones,
                        precio_unitario=precio,
                        iva_pct=iva_pct,
                        aplica_descuento_escala=aplica_escala,
                        activo=True,
                    ))
                    creados += 1

            except Exception as e:
                errores.append(f"{fila}: {e}")

        # Cargar escalas si la hoja existe
        escalas_creadas = 0
        if "Escalas" in wb.sheetnames:
            hoja_esc = wb["Escalas"]
            enc_esc = [str(c.value).strip().lower() if c.value else "" for c in hoja_esc[1]]
            idx_esc = {nombre: i for i, nombre in enumerate(enc_esc)}

            # Reemplazar todas las escalas
            db.query(EscalaDescuento).delete()
            for fila in hoja_esc.iter_rows(min_row=2, values_only=True):
                if not fila or all(v is None for v in fila):
                    continue
                try:
                    cat_esc = None
                    if "categoria" in idx_esc and fila[idx_esc["categoria"]]:
                        cat_esc = str(fila[idx_esc["categoria"]]).strip()

                    db.add(EscalaDescuento(
                        categoria=cat_esc,
                        escala=int(fila[idx_esc["escala"]]),
                        cantidad_min=int(fila[idx_esc["cantidad_min"]]),
                        cantidad_max=(
                            int(fila[idx_esc["cantidad_max"]])
                            if "cantidad_max" in idx_esc and fila[idx_esc["cantidad_max"]] is not None
                            else None
                        ),
                        descuento_pct=float(fila[idx_esc["descuento_pct"]]),
                    ))
                    escalas_creadas += 1
                except Exception as e:
                    errores.append(f"Escala: {e}")

        db.commit()
        print(f"\n✅ RESULTADO:")
        print(f"   Productos creados:      {creados}")
        print(f"   Productos actualizados: {actualizados}")
        print(f"   Escalas cargadas:       {escalas_creadas}")
        if errores:
            print(f"\n⚠️  {len(errores)} errores:")
            for e in errores[:10]:
                print(f"   - {e}")

    except Exception as e:
        db.rollback()
        print(f"❌ Error: {e}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    cargar()